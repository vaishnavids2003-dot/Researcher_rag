
from pathlib import Path

from pypdf import PdfReader
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx"}


def load_document(file_path: str) -> list[dict]:
    """Extract text and metadata from PDF, TXT, or DOCX files."""

    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"File not found: {file_path}")

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file format: {extension}")

    if extension == ".pdf":
        documents = _load_pdf(path)
    elif extension == ".txt":
        documents = _load_txt(path)
    else:
        documents = _load_docx(path)

    if not any(item["text"].strip() for item in documents):
        raise ValueError(
            f"No extractable text found in {path.name}"
        )

    return documents


def _load_pdf(path: Path) -> list[dict]:
    """Extract text page by page using pypdf."""

    documents = []
    reader = PdfReader(str(path))

    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()

        if text:
            documents.append(
                {
                    "text": text,
                    "metadata": {
                        "source": path.name,
                        "page": page_number,
                    },
                }
            )

    return documents


def _load_txt(path: Path) -> list[dict]:
    """Read a UTF-8 text file."""

    text = path.read_text(encoding="utf-8-sig").strip()

    return [
        {
            "text": text,
            "metadata": {"source": path.name},
        }
    ]


def _load_docx(path: Path) -> list[dict]:
    """Extract paragraphs and tables from a DOCX file."""

    document = Document(str(path))
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]

            if any(cells):
                parts.append(" | ".join(cells))

    full_text = "\n".join(parts).strip()

    return [
        {
            "text": full_text,
            "metadata": {"source": path.name},
        }
    ]
